"""Auction tests use synthetic workbooks and isolated temporary storage only."""
import io
import tempfile
from datetime import timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.sessions.models import Session
from django.utils import timezone
from django.test import Client, SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from ..auction_forms import AuctionFilterForm
from ..services import auction, auction_storage
from ..views.auction import SESSION_STATE, initial_state
from .support import ACCOUNT_PASSWORD, TEST_ENCRYPTION_KEY


def excel(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class AuctionServiceTests(SimpleTestCase):
    def test_aliases_boundaries_and_invalid_values_partition_without_changing_data(self):
        values = [[62, 48000], [63, 47000], [61, 48001], [None, 1], ['bad', 1],
                  [-1, 1], [0, 1], [True, 1], ['inf', 1], [60, 'NaN'], ['62', '48000']]
        table = {'columns': ['WIDTH (inches)', 'Weight Pounds'], 'rows': values}
        groups = auction.partition(table, auction.DEFAULT_FILTERS)
        self.assertEqual(groups['eligible'], [values[0], values[-1]])
        self.assertEqual(groups['excluded'], values[1:-1])
        self.assertEqual(table['rows'], values)

    def test_all_comparisons_combine_with_and(self):
        table = {'columns': ['Width', 'Weight lb'], 'rows': [[60, 100], [62, 200], [64, 300], [0, 400]]}
        for operator, expected in [('eq', [[62, 200]]), ('le', [[60, 100], [62, 200]]), ('ge', [[62, 200], [64, 300]])]:
            filters = {'width': '62', 'weight': '200', 'width_op': operator, 'weight_op': operator}
            self.assertEqual(auction.partition(table, filters)['eligible'], expected)
        filters.update(width_op='le', weight_op='ge')
        self.assertEqual(auction.partition(table, filters)['eligible'], [[62, 200]])

    def test_filter_validation(self):
        for value in ('NaN', 'Infinity', '-1', '0', 'text'):
            self.assertFalse(AuctionFilterForm({**auction.DEFAULT_FILTERS, 'width': value}).is_valid())
        self.assertFalse(AuctionFilterForm({**auction.DEFAULT_FILTERS, 'weight_op': 'gt'}).is_valid())

    def test_csv_preserves_strings_and_first_row_as_headers(self):
        table = auction.read_table('test.csv', b'\xef\xbb\xbfWidth,Weight lbs,Tag\r\n62,48000,0012\r\n')
        self.assertEqual(table['rows'], [['62', '48000', '0012']])
        self.assertEqual(auction.inspect_upload('test.CSV', b'file'), [])

    def test_excel_sheet_selection_and_values(self):
        content = excel({'First': [['Width', 'Weight lbs'], [60, 100]], 'Second': [['Width', 'Weight lbs'], [62, 48000]]})
        self.assertEqual(auction.inspect_upload('test.xlsx', content), ['First', 'Second'])
        self.assertEqual(auction.read_table('test.xlsx', content, 'Second')['rows'], [[62, 48000]])
        with self.assertRaises(auction.AuctionError):
            auction.read_table('test.xlsx', content, 'Unknown')

    def test_invalid_files_headers_and_limits(self):
        for name, content in [('file.txt', b'x'), ('file.xlsx', b'not an excel'), ('file.csv', b'')]:
            with self.subTest(name=name), self.assertRaises(auction.AuctionError):
                auction.inspect_upload(name, content)
        for content in (b'', b'Width,Weight kg\n1,2', b'Width,Width,Weight lbs\n1,1,2',
                        b'Width,Weight lbs\n1,2,3', b'Width,Weight lbs\n\xff,2', b'Width,Weight lbs\n"unterminated,2'):
            with self.subTest(content=content), self.assertRaises(auction.AuctionError):
                auction.read_table('test.csv', content)
        with patch.object(auction, 'MAX_ROWS', 1), self.assertRaises(auction.AuctionError):
            auction.read_table('test.csv', b'Width,Weight lbs\n1,2\n3,4')
        with patch.object(auction, 'MAX_FILE_BYTES', 1), self.assertRaises(auction.AuctionError):
            auction.inspect_upload('test.csv', b'12')
        with patch.object(auction, 'MAX_EXPANDED_BYTES', 1), self.assertRaises(auction.AuctionError):
            auction.inspect_upload('test.xlsx', excel({'Sheet': [['Width', 'Weight lbs']]}))

    def test_sort_numeric_text_missing_and_restore(self):
        rows = [['10'], ['2'], [None], [''], ['apple'], ['Banana']]
        self.assertEqual(auction.sorted_rows(rows, 0, 'asc'), [['2'], ['10'], ['apple'], ['Banana'], [None], ['']])
        self.assertEqual(auction.sorted_rows(rows, 0, 'desc'), [['Banana'], ['apple'], ['10'], ['2'], [None], ['']])
        self.assertEqual(auction.sorted_rows(rows), rows)
        self.assertEqual(auction.sorted_rows([['alpha', 1], ['Alpha', 2]], 0, 'desc'), [['alpha', 1], ['Alpha', 2]])

    def test_pagination_does_not_cap_at_history_limit(self):
        state = initial_state('token')
        state['pages']['eligible'] = 11
        page = auction.page_table(['ID'], [[n] for n in range(1005)], state, 'eligible')
        self.assertEqual(page['rows'], [[n] for n in range(1000, 1005)])
        self.assertEqual((page['page'], page['pages']), (11, 11))
        empty = auction.page_table(['ID'], [], state, 'excluded')
        self.assertEqual((empty['start'], empty['end'], empty['page']), (0, 0, 1))

    def test_export_includes_both_complete_sorted_sets_and_literal_formula_text(self):
        table = {'columns': ['Width', 'Weight lbs', '=HEADER'], 'rows': []}
        groups = {'eligible': [[60, 100, '=1+1'] for _ in range(105)], 'excluded': [[63, 100, '+SUM(A1:A2)']]}
        content = auction.export_workbook(table, groups, initial_state('token')['sorts'])
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, ['Eligible', 'Excluded'])
        self.assertEqual(workbook['Eligible'].max_row, 106)
        self.assertEqual(workbook['Excluded'].max_row, 2)
        self.assertEqual(workbook['Eligible']['C2'].data_type, 's')
        self.assertEqual(workbook['Eligible']['C1'].data_type, 's')
        self.assertEqual(workbook['Eligible']['C2'].value, '=1+1')
        workbook.close()


@override_settings(CREDENTIAL_ENCRYPTION_KEY=TEST_ENCRYPTION_KEY)
class AuctionWorkflowTests(TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        override = override_settings(AUCTION_UPLOAD_DIR=Path(directory.name))
        override.enable()
        self.addCleanup(override.disable)
        self.directory = Path(directory.name)
        self.user = User.objects.create_user('auction-buyer', password=ACCOUNT_PASSWORD)
        self.client.force_login(self.user)

    def upload(self, content=b'Width,Weight lbs,Tag\n62,48000,good\n63,47000,bad\n', name='synthetic.csv'):
        return self.client.post(reverse('auction'), {'action': 'upload', 'file': SimpleUploadedFile(name, content)}, follow=True)

    def action(self, action, **fields):
        return self.client.post(reverse('auction'), {'action': action, 'token': self.client.session[SESSION_STATE]['token'], **fields}, follow=True)

    def test_empty_page_and_authentication(self):
        response = self.client.get(reverse('auction'))
        self.assertContains(response, 'Upload Auction List')
        self.assertNotContains(response, 'Eligible Products')
        self.assertNotContains(response, 'Excluded Products')
        for name in ('auction', 'auction-download'):
            self.assertEqual(Client().get(reverse(name)).status_code, 302)

    def test_upload_defaults_show_hide_and_filters_are_synchronized(self):
        response = self.upload()
        self.assertContains(response, 'Filtered out 1 rows.')
        self.assertNotContains(response, 'id="excluded-card"')
        response = self.action('show')
        self.assertContains(response, 'id="excluded-card"')
        self.assertContains(response, 'Hide Excluded Table')
        response = self.action('filter', width_op='ge', width='63', weight_op='le', weight='48000')
        self.assertEqual(response.context['eligible']['rows'][0][2], 'bad')
        self.assertEqual(response.context['excluded']['rows'][0][2], 'good')
        self.assertEqual(response.context['eligible_filter'].initial, response.context['excluded_filter'].initial)
        response = self.action('hide')
        self.assertNotContains(response, 'id="excluded-card"')

    def test_multi_sheet_requires_selection_and_changes_reset_state(self):
        content = excel({'One': [['Width', 'Weight lbs'], [60, 100]], 'Two': [['Width', 'Weight lbs'], [63, 200]]})
        response = self.upload(content, 'two.xlsx')
        self.assertContains(response, 'Choose a worksheet')
        self.assertNotContains(response, 'Eligible Products')
        self.action('sheet', sheet='One')
        self.action('show')
        self.action('sort', table='eligible', column=0)
        self.action('filter', **{**auction.DEFAULT_FILTERS, 'width': '100'})
        response = self.action('sheet', sheet='Two')
        state = self.client.session[SESSION_STATE]
        self.assertEqual(state['filters'], auction.DEFAULT_FILTERS)
        self.assertFalse(state['show_excluded'])
        self.assertIsNone(state['sorts']['eligible']['column'])
        self.assertEqual(response.context['excluded']['total'], 1)

    def test_single_excel_automatically_loads(self):
        response = self.upload(excel({'Only': [['Width', 'Weight lbs'], [62, 48000]]}), 'single.xlsx')
        self.assertEqual(response.context['eligible']['total'], 1)
        self.assertNotContains(response, 'Choose a worksheet')

    def test_invalid_sheet_clears_previous_results(self):
        self.upload(excel({'Good': [['Width', 'Weight lbs'], [60, 100]], 'Bad': [['Other'], [1]]}), 'multi.xlsx')
        self.action('sheet', sheet='Good')
        response = self.action('sheet', sheet='Bad')
        self.assertEqual(response.status_code, 400)
        self.assertNotContains(response, 'Eligible Products', status_code=400)
        self.assertEqual(self.client.get(reverse('auction-download')).status_code, 400)

    def test_invalid_replacement_clears_previous_upload_and_successful_replacement_resets(self):
        self.upload()
        self.action('show')
        self.action('sort', table='eligible', column=0)
        self.upload(b'Width,Weight lbs\n60,100', 'replacement.csv')
        self.assertEqual(len(list(self.directory.glob('*.json'))), 1)
        self.assertFalse(self.client.session[SESSION_STATE]['show_excluded'])
        response = self.upload(b'bad', 'bad.txt')
        self.assertEqual(response.status_code, 400)
        self.assertNotIn(SESSION_STATE, self.client.session)
        self.assertEqual(len(list(self.directory.glob('*.json'))), 0)

    def test_pages_and_sort_cycle_are_independent_and_reset_on_apply(self):
        rows = ['Width,Weight lbs,Tag'] + [f'60,{1000-i},E{i}' for i in range(205)] + [f'70,{1000-i},X{i}' for i in range(105)]
        self.upload('\n'.join(rows).encode())
        response = self.action('paginate', table='eligible', page=3)
        self.assertEqual(len(response.context['eligible']['rows']), 5)
        self.assertEqual(response.context['excluded']['page'], 1)
        response = self.action('sort', table='eligible', column=1)
        self.assertEqual(response.context['eligible']['rows'][0][1], '796')
        self.assertEqual(response.context['eligible']['page'], 1)
        response = self.action('sort', table='eligible', column=1)
        self.assertEqual(response.context['eligible']['rows'][0][1], '1000')
        self.action('sort', table='eligible', column=1)
        self.assertIsNone(self.client.session[SESSION_STATE]['sorts']['eligible']['column'])
        self.action('paginate', table='excluded', page=2)
        self.action('filter', **auction.DEFAULT_FILTERS)
        self.assertEqual(self.client.session[SESSION_STATE]['pages'], {'eligible': 1, 'excluded': 1})

    def test_invalid_controls_leave_applied_filters_unchanged(self):
        self.upload()
        for action, fields in [('filter', {**auction.DEFAULT_FILTERS, 'width': 'NaN'}),
                               ('paginate', {'table': 'eligible', 'page': 99}),
                               ('sort', {'table': 'eligible', 'column': 99})]:
            self.assertEqual(self.action(action, **fields).status_code, 400)
        self.assertEqual(self.client.session[SESSION_STATE]['filters'], auction.DEFAULT_FILTERS)

    def test_export_uses_applied_filters_and_includes_hidden_excluded(self):
        self.upload()
        self.action('filter', width='62', width_op='ge', weight='48000', weight_op='eq')
        response = self.client.get(reverse('auction-download'), {'width': '100'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook['Eligible'].max_row, 2)
        self.assertEqual(workbook['Excluded'].max_row, 2)
        self.assertEqual(workbook['Eligible']['C2'].value, 'good')
        workbook.close()

    def test_session_and_user_isolation_and_cleanup_on_next_login(self):
        self.upload()
        state = self.client.session[SESSION_STATE]
        for user in (self.user, User.objects.create_user('other-buyer', password=ACCOUNT_PASSWORD)):
            other = Client()
            other.force_login(user)
            session = other.session
            session[SESSION_STATE] = state
            session.save()
            response = other.get(reverse('auction'))
            self.assertEqual(response.status_code, 400)
            self.assertNotIn(SESSION_STATE, other.session)
        self.assertEqual(self.client.get(reverse('auction-download')).status_code, 200)
        self.client.post(reverse('logout'))
        self.assertEqual(len(list(self.directory.glob('*.json'))), 1)
        self.assertEqual(self.client.get(reverse('auction-download')).status_code, 302)
        self.client.post(reverse('login'), {'username': self.user.username, 'password': ACCOUNT_PASSWORD})
        self.assertEqual(list(self.directory.glob('*.json')), [])
        self.assertNotIn(SESSION_STATE, self.client.session)

    def test_expired_session_is_cleaned_only_after_successful_login(self):
        self.upload()
        file = next(self.directory.glob('*.json'))
        self.assertEqual(file.stat().st_mode & 0o777, 0o600)
        Session.objects.filter(session_key=self.client.session.session_key).update(
            expire_date=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.client.get(reverse('auction')).status_code, 302)
        self.assertTrue(file.exists())
        self.client.post(reverse('login'), {'username': self.user.username, 'password': 'wrong'})
        self.assertTrue(file.exists())
        response = self.client.post(reverse('login'), {'username': self.user.username, 'password': ACCOUNT_PASSWORD})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(file.exists())
        self.assertNotIn(SESSION_STATE, self.client.session)

    def test_login_preserves_other_active_sessions_and_other_users_uploads(self):
        self.upload()
        first = next(self.directory.glob('*.json'))
        other_user = User.objects.create_user('separate-buyer', password=ACCOUNT_PASSWORD)
        other = Client()
        other.force_login(other_user)
        other.post(reverse('auction'), {'action': 'upload', 'file': SimpleUploadedFile('other.csv', b'Width,Weight lbs\n60,100')})
        other.post(reverse('logout'))
        second = next(path for path in self.directory.glob('*.json') if path != first)
        fresh = Client()
        fresh.post(reverse('login'), {'username': self.user.username, 'password': ACCOUNT_PASSWORD})
        self.assertTrue(first.exists())
        self.assertTrue(second.exists())
        self.assertEqual(self.client.get(reverse('auction-download')).status_code, 200)

    def test_active_upload_has_no_separate_expiration_timer(self):
        self.upload()
        state = self.client.session[SESSION_STATE]
        snapshot = auction_storage.load(state['token'], self.user.pk, self.client.session.session_key)
        self.assertNotIn('expires', snapshot)
        # Older snapshots may carry an upload deadline; the login session now controls access.
        snapshot['expires'] = 0
        auction_storage.save(state['token'], snapshot)
        self.assertEqual(self.client.get(reverse('auction-download')).status_code, 200)

    def test_csrf_and_escaped_cell_contents(self):
        protected = Client(enforce_csrf_checks=True)
        protected.force_login(self.user)
        self.assertEqual(protected.post(reverse('auction'), {'action': 'upload'}).status_code, 403)
        response = self.upload(b'Width,Weight lbs,Tag\n60,100,<script>alert(1)</script>')
        self.assertContains(response, '&lt;script&gt;alert(1)&lt;/script&gt;')
        self.assertNotContains(response, '<script>alert(1)</script>')
        self.assertIn('no-store', response['Cache-Control'])
