"""Auction spreadsheet parsing, screening, sorting, pagination, and Excel export.

These functions operate on plain tables and never query the client database.
"""
import csv
import io
import math
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLUMNS = 100
PAGE_SIZE = 100
DEFAULT_FILTERS = {'width_op': 'le', 'width': '62', 'weight_op': 'le', 'weight': '48000'}
OPERATORS = {'eq': lambda a, b: a == b, 'le': lambda a, b: a <= b, 'ge': lambda a, b: a >= b}


class AuctionError(ValueError):
    """An upload or requested table operation cannot be processed."""


def _workbook(content):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise AuctionError('This workbook is too large after decompression. Use a smaller file.')
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except AuctionError:
        raise
    except Exception as error:
        raise AuctionError('Unable to read this workbook. Upload a valid, unencrypted .xlsx file.') from error


def inspect_upload(name, content):
    """Validate the file and return sheet names, without choosing a multi-sheet workbook."""
    if not content or len(content) > MAX_FILE_BYTES:
        raise AuctionError('Choose a nonempty file no larger than 10 MB.')
    extension = Path(name).suffix.lower()
    if extension == '.csv':
        return []
    if extension != '.xlsx':
        raise AuctionError('Choose an .xlsx or .csv file.')
    workbook = _workbook(content)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def measurement_columns(columns):
    """Match the notebook's aliases while retaining original column labels."""
    normalized = [re.sub(r'[^a-z0-9]', '', str(column).lower()) for column in columns]
    widths = [i for i, name in enumerate(normalized) if name in {'width', 'widthin', 'widthinch', 'widthinches'}]
    weights = [i for i, name in enumerate(normalized) if name in {'weightlbs', 'weightlb', 'weightpounds'}]
    if len(widths) != 1 or len(weights) != 1:
        raise AuctionError('Expected one Width column and one Weight lbs column. Check the column names.')
    return widths[0], weights[0]


def _table(records):
    """Validate headers and row limits without silently truncating an upload."""
    header = next(records, None)
    if header is None:
        raise AuctionError('The selected sheet is empty.')
    columns = [str(value) if value is not None else '' for value in header]
    if not columns or len(columns) > MAX_COLUMNS:
        raise AuctionError('Use a sheet with no more than 100 columns.')
    if any(not value.strip() or len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value) for value in columns) or len(set(columns)) != len(columns):
        raise AuctionError('Each column must have a nonempty, unique heading in the first row.')
    measurement_columns(columns)
    rows = []
    for row in records:
        if len(rows) >= MAX_ROWS:
            raise AuctionError('Use a sheet with no more than 100,000 data rows.')
        if len(row) != len(columns):
            raise AuctionError('Every row must have the same number of columns as the header.')
        values = [value.isoformat() if isinstance(value, (date, datetime, time)) else value for value in row]
        if any(isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)) for value in values):
            raise AuctionError('A cell contains text that Excel cannot export. Remove control characters or shorten it to 32,767 characters.')
        rows.append([str(value) if isinstance(value, float) and not math.isfinite(value) else value for value in values])
    return {'columns': columns, 'rows': rows}


def read_table(name, content, sheet=None):
    """Read UTF-8 CSV or the explicitly selected Excel worksheet."""
    if Path(name).suffix.lower() == '.csv':
        try:
            text = content.decode('utf-8-sig')
            return _table(iter(csv.reader(io.StringIO(text), strict=True)))
        except (UnicodeError, csv.Error) as error:
            raise AuctionError('Unable to read this CSV. Use a UTF-8, comma-separated file.') from error
    workbook = _workbook(content)
    try:
        if sheet not in workbook.sheetnames:
            raise AuctionError('Select a worksheet from this workbook.')
        worksheet = workbook[sheet]
        if (worksheet.max_column or 0) > MAX_COLUMNS or (worksheet.max_row or 0) > MAX_ROWS + 1:
            raise AuctionError('Use a sheet with at most 100 columns and 100,000 data rows.')
        # Formulas remain literal text; they are never executed or treated as measurements.
        return _table(iter(worksheet.iter_rows(values_only=True)))
    except AuctionError:
        raise
    except Exception as error:
        raise AuctionError('Unable to read this worksheet. Check its contents and try again.') from error
    finally:
        workbook.close()


def number(value):
    """Convert comparisons only; original cell values stay intact for display/export."""
    if isinstance(value, bool) or value is None or str(value).strip() == '':
        return None
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def partition(table, filters):
    """Every row appears exactly once; invalid measurements are always excluded."""
    width_index, weight_index = measurement_columns(table['columns'])
    limits = [number(filters['width']), number(filters['weight'])]
    if any(value is None or value <= 0 for value in limits):
        raise AuctionError('Width and weight must be positive, finite numbers.')
    try:
        comparisons = [OPERATORS[filters['width_op']], OPERATORS[filters['weight_op']]]
    except KeyError as error:
        raise AuctionError('Choose =, ≤, or ≥ for each filter.') from error
    eligible, excluded = [], []
    for row in table['rows']:
        values = [number(row[width_index]), number(row[weight_index])]
        valid = all(value is not None and value > 0 for value in values)
        fits = valid and all(compare(value, limit) for compare, value, limit in zip(comparisons, values, limits))
        (eligible if fits else excluded).append(row)
    return {'eligible': eligible, 'excluded': excluded}


def sorted_rows(rows, column=None, direction=''):
    """Stable sorting across all rows, with blanks last in either direction."""
    if column is None:
        return rows
    present = [row for row in rows if row[column] is not None and row[column] != '']
    missing = [row for row in rows if row[column] is None or row[column] == '']
    def key(row):
        value = row[column]
        numeric = number(value)
        return (0, numeric) if numeric is not None else (1, str(value).casefold())
    return sorted(present, key=key, reverse=direction == 'desc') + missing


def page_table(columns, rows, state, kind):
    order = state['sorts'][kind]
    rows = sorted_rows(rows, order['column'], order['direction'])
    count = len(rows)
    pages = max(1, math.ceil(count / PAGE_SIZE))
    page = min(max(1, state['pages'][kind]), pages)
    start = (page - 1) * PAGE_SIZE
    headers = []
    for index, column in enumerate(columns):
        active = index == order['column']
        direction = order['direction'] if active else ''
        headers.append({'label': column, 'index': index,
                        'aria_sort': {'asc': 'ascending', 'desc': 'descending'}.get(direction, 'none'),
                        'icon': {'asc': '↑', 'desc': '↓'}.get(direction, '↕'),
                        'next_label': {'asc': 'Sort descending', 'desc': 'Restore original order'}.get(direction, 'Sort ascending')})
    return {'kind': kind, 'headers': headers, 'rows': rows[start:start + PAGE_SIZE],
            'total': count, 'page': page, 'pages': pages, 'previous': max(1, page - 1),
            'next': min(pages, page + 1), 'start': start + 1 if count else 0,
            'end': min(start + PAGE_SIZE, count)}


def export_workbook(table, groups, sorts):
    """Export all rows in each group's applied order; force strings to text cells."""
    workbook = Workbook(write_only=True)
    for kind in ('eligible', 'excluded'):
        sheet = workbook.create_sheet(kind.title())
        sheet.freeze_panes = 'A2'
        order = sorts[kind]
        rows = sorted_rows(groups[kind], order['column'], order['direction'])
        for row in [table['columns']]:
            _append_text_safe(sheet, row)
        for row in rows:
            _append_text_safe(sheet, row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _append_text_safe(sheet, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = 's'
        cells.append(cell)
    sheet.append(cells)
